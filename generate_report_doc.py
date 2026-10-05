import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = docx.Document()

    # 페이지 여백 설정 (기본 1인치)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # 기본 스타일 설정
    style_normal = doc.styles['Normal']
    style_normal.font.name = '맑은 고딕'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # 1. 문서 제목 (Title)
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = title_p.add_run("2026학년도 캡스톤디자인 최종 결과보고서\n")
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x00, 0x3B, 0x70)

    run_title = title_p.add_run("남서울대학교 지능형 캠퍼스 라이프 종합 플랫폼\n(Chat-NSU & Smart Campus Platform)")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    # 구분선 대체 여백
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(20)
    run_meta = p_meta.add_run("플랫폼 배포 URL: https://singular-moonbeam-adafa6.netlify.app | 기술 스택: Flutter & Google Gemini AI")
    run_meta.font.size = Pt(9.5)
    run_meta.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # 정보 표 (과제명, 팀원 등)
    info_table = doc.add_table(rows=4, cols=4)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False

    table_data = [
        [("과제명", True), ("남서울대 지능형 캠퍼스 라이프 및 포트폴리오 통합 앱", False), ("개발구분", True), ("웹 & 모바일 크로스플랫폼", False)],
        [("소속대학", True), ("남서울대학교 공과대학", False), ("학과/전공", True), ("컴퓨터소프트웨어 / 지능정보통신", False)],
        [("개발기간", True), ("2026. 03 ~ 2026. 10", False), ("라이브 배포", True), ("Netlify Production Live", False)],
        [("핵심기술", True), ("Flutter, Dart, Gemini AI API, Headless Web Parsing, Netlify", False, 3)]
    ]

    for r_idx, row_info in enumerate(table_data):
        row = info_table.rows[r_idx]
        c_idx = 0
        for item in row_info:
            text = item[0]
            is_header = item[1]
            colspan = item[2] if len(item) > 2 else 1
            cell = row.cells[c_idx]
            cell.text = text
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            if is_header:
                set_cell_background(cell, "F0F4F8")
                cell.paragraphs[0].runs[0].font.bold = True
                cell.paragraphs[0].runs[0].font.size = Pt(9.5)
                cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x00, 0x3B, 0x70)
            else:
                cell.paragraphs[0].runs[0].font.size = Pt(9.5)
            
            if colspan > 1:
                # 병합 처리
                cell.merge(row.cells[c_idx + colspan - 1])
                c_idx += colspan
            else:
                c_idx += 1

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x00, 0x3B, 0x70)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        return p

    def add_body(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(4)
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.font.bold = True
        p.add_run(text)
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(3)
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.font.bold = True
        p.add_run(text)
        return p

    # 1. 서론 및 프로젝트 개요
    add_h1("1. 프로젝트 개요 (Introduction)")
    add_h2("1.1 개발 배경 및 필요성")
    add_body("대학 캠퍼스 생활에서 학생들은 학사일정, 강의실 위치, 셔틀버스 시간, 학식 메뉴, 도서관 좌석 현황 등 방대한 정보에 수시로 접근해야 합니다. 그러나 남서울대학교의 정보 시스템은 포털 사이트, 도서관 좌석 시스템, 학생회관 식단 공지 등이 각기 다른 웹페이지로 분절되어 있어 재학생들이 일상 정보를 신속하게 확인하는 데 큰 불편함을 겪고 있습니다.")
    add_body("특히 기존 학교 식당 안내의 경우 정적인 웹 구조가 아닌 동적 AJAX 게시판 기반으로 제공되어 모바일 브라우저 접근성이 떨어지고, 70여 개가 넘는 코너별 메뉴와 가격표를 한눈에 파악하기 어려운 문제가 있었습니다. 또한 대학 생활 중 축적되는 자격증, 프로젝트, 수상 실적을 취업용 포트폴리오로 체계화할 수 있는 통합 도구의 부재도 중요한 개선 과제였습니다.")

    add_h2("1.2 프로젝트 목표 및 핵심 가치")
    add_bullet(" 재학생을 위한 원스톱 모바일/웹 캠퍼스 라이프 플랫폼 구축", bold_prefix="통합 정보 허브: ")
    add_bullet(" 남서울대 학칙, 수강신청, 졸업요건, 장학제도 등을 즉시 답변하는 AI 도우미 탑재", bold_prefix="지능형 AI 챗봇: ")
    add_bullet(" 학교 공식 CMS를 헤드리스 파싱하여 100% 무결한 실제 식당 메뉴 및 가격표 제공", bold_prefix="공식 학식 무결성: ")
    add_bullet(" 성암기념중앙도서관 실시간 좌석 현황 파싱 및 강의실/편의시설 가이드 구축", bold_prefix="실시간 편의 기능: ")
    add_bullet(" 학기별 프로젝트, 기술 스택, 수상 경력을 정리할 수 있는 표준 스펙 구현", bold_prefix="커리어 포트폴리오: ")

    # 2. 시스템 아키텍처 및 기술 스택
    add_h1("2. 시스템 아키텍처 및 기술 스택 (System Architecture)")
    add_body("본 프로젝트는 반응형 모바일 애플리케이션과 웹 브라우저 환경을 동시에 지원하는 단일 코드베이스(Cross-Platform) 구조로 설계되었습니다.")

    tech_table = doc.add_table(rows=6, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_table.autofit = False

    t_headers = ["구분", "적용 기술 / 라이브러리", "담당 역할 및 특징"]
    for i, h in enumerate(t_headers):
        cell = tech_table.rows[0].cells[i]
        cell.text = h
        set_cell_background(cell, "003B70")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.paragraphs[0].runs[0].font.size = Pt(9.5)

    t_rows = [
        ("Frontend", "Flutter 3.x, Dart 3.x", "Material 3 디자인, 반응형 멀티플랫폼 UI, 모달 바텀시트"),
        ("AI Engine", "Google Gemini AI API", "남서울대학교 학사/편의 데이터셋 프롬프트 엔지니어링"),
        ("Data Scraping", "Python, Edge Headless DOM Crawler", "학교 CMS AJAX 동적 렌더링 테이블(Board 465, 466, 467) 전수 파싱"),
        ("Library Seat API", "HTTP Client, Regex/HTML Parser", "성암기념중앙도서관 열람실 실시간 잔여석 파싱 및 시각화"),
        ("Deployment", "GitHub, Netlify Cloud Hosting", "릴리즈 웹 빌드 자동 동기화 및 24시간 실시간 무중단 호스팅")
    ]

    for r_idx, r_data in enumerate(t_rows, start=1):
        row = tech_table.rows[r_idx]
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.text = val
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            if r_idx % 2 == 1:
                set_cell_background(cell, "F9FAFC")
            cell.paragraphs[0].runs[0].font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3. 핵심 기능 및 상세 구현 내용
    add_h1("3. 핵심 기능 및 상세 구현 내용 (Key Features)")

    add_h2("3.1 100% 공식 데이터 연동 학식 & 푸드코트 시스템")
    add_body("본 프로젝트의 핵심 성과 중 하나는 학교 공식 홈페이지에 게시된 실제 키오스크 등록 메뉴 및 가격표를 한치의 오차 없이 100% 반영한 것입니다.")
    add_bullet(" 수제버거 브랜드 '버거운베어'(16종), 덮밥·라면 브랜드 '청년한끼'(23종), '쑝쑝돈까스'(16종), 백반 브랜드 '태산김치찜'(16종) 총 71종의 메뉴명, 단품/세트 구성 및 공식 가격 완벽 연동", bold_prefix="학생회관 1층 푸드코트 (71종): ")
    add_bullet(" 한식국밥 '언덕뜰소반'(16종), 직화찌개 '부대통령'(16종), 쌀국수 '포한끼'(16종), '한우사골마라탕'(16종) 총 64종의 고정 정규 메뉴 완벽 연동", bold_prefix="학생회관 2층 학생식당 (64종): ")
    add_bullet(" 학생 및 교직원 누구나 이용 가능한 6,500원 자율배식 한식 뷔페 식단표(제육볶음, 들기름막국수 등) 및 요일별 영양 식단 안내", bold_prefix="학생회관 3층 교직원식당: ")
    add_bullet(" 재학생 대상 1,000원 천원의 아침밥(선착순 100명) 및 기숙사 중·석식 뷔페 및 상시 식당 메뉴 제공", bold_prefix="엘림생활관 2관 멀베리: ")

    add_h2("3.2 사용자 중심 UI/UX 최적화: 컴팩트 뷰 & 바텀시트 모달")
    add_body("총 135종 이상의 방대한 메뉴를 메인 화면에 한꺼번에 나열할 경우, 스크롤이 지나치게 길어져 다른 캠퍼스 편의 기능 접근을 방해하는 치명적인 UI 문제가 있었습니다. 이를 해결하기 위해 다음과 같이 2단계 UI 구조로 리팩토링을 완료했습니다:")
    add_bullet(" 각 층 선택 칩(1층/2층/3층/멀베리) 터치 시, 운영시간 안내와 함께 코너별 브랜드 태그, 4개 코너 대표 메뉴 샘플만 콤팩트하게 노출하여 메인 카드의 세로 높이를 대폭 축소", bold_prefix="메인 화면 요약 카드: ")
    add_bullet(" 각 층 하단에 '전체 메뉴 & 가격표 전체보기 📋' 버튼을 배치하여, 터치 시 부드럽게 열리는 DraggableScrollableSheet 모달 팝업 구현", bold_prefix="전체보기 모달 (BottomSheet): ")
    add_bullet(" 모달 상단에 '전체' 및 각 입점 코너별 필터 칩을 제공하여 학생들이 원하는 식당의 메뉴만 골라보거나 가격대별로 손쉽게 탐색 가능하도록 편의성 극대화", bold_prefix="코너별 인터랙티브 필터: ")

    add_h2("3.3 데이터 순수성 및 신뢰성 보장 (No False Data)")
    add_body("사용자 피드백을 적극 수용하여, 실시간 판매 순위나 공식 통계가 아닌 개발자 임의의 '⭐ 인기' 또는 '대표' 뱃지를 전면 삭제하였습니다. 이를 통해 학교 공식 홈페이지에 고시된 원본 데이터와 100% 일치하는 객관적이고 신뢰할 수 있는 정보를 제공합니다.")

    add_h2("3.4 지능형 캠퍼스 AI 챗봇 시스템")
    add_body("남서울대학교 학사 일정, 수강신청 가이드, 졸업 요건(외국어 성적 및 캡스톤 이수 등), 통학 셔틀버스 노선 및 승차장 위치, 교내 장학금 기준 등 재학생들이 가장 많이 묻는 교내 필수 질문에 대해 Google Gemini API 기반 지능형 답변을 즉각 생성합니다.")

    add_h2("3.5 실시간 중앙도서관 좌석 현황 & 캠퍼스 맵 네비게이션")
    add_body("성암기념중앙도서관의 일반열람실, 노트북열람실 실시간 잔여석 데이터를 실시간 파싱하여 도서관 방문 전 혼잡도를 미리 파악할 수 있습니다. 또한 공학1관, 공학2관, 학생회관, 성암문화관 등 주요 호관별 층별 안내도 및 복사실, 보건실, 무인발급기 등 공식 편의시설 위치를 직관적으로 제공합니다.")

    add_h2("3.6 재학생 커리어 포트폴리오 관리 시스템")
    add_body("재학 중 진행한 캡스톤 프로젝트, 동아리 활동, 취득 자격증, 어학 점수를 체계적으로 등록하고 취업 준비 시 이력서 포맷으로 출력 및 관리할 수 있는 맞춤형 포트폴리오 데이터 스펙을 구현했습니다.")

    # 4. 문제 해결 및 리팩토링 과정
    add_h1("4. 문제 해결 및 리팩토링 과정 (Troubleshooting)")
    
    trouble_table = doc.add_table(rows=5, cols=3)
    trouble_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    trouble_table.autofit = False

    t2_headers = ["발생 문제 (Issue)", "원인 분석 (Root Cause)", "해결 방안 및 최종 성과 (Solution & Result)"]
    for i, h in enumerate(t2_headers):
        cell = trouble_table.rows[0].cells[i]
        cell.text = h
        set_cell_background(cell, "003B70")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.paragraphs[0].runs[0].font.size = Pt(9.5)

    t2_rows = [
        ("학교 홈페이지 학식 데이터 빈 화면 현상", "학교 CMS가 AJAX 동적 스크립트로 테이블을 로드하여 단순 HTML 요청 시 테이블이 비어있음", "Headless Edge 브라우저 자동화 스크립트를 작성하여 실제 DOM을 렌더링한 후 Board 465, 466, 467 전체 메뉴 135종 추출 성공"),
        ("메뉴 과다로 인한 모바일 스크롤 과부하", "70여 개 메뉴를 메인 카드에 수직으로 일괄 렌더링하여 화면 스크롤이 수천 픽셀로 과도하게 늘어남", "메인 화면은 4개 코너 요약 및 샘플만 노출하고, 전용 바텀시트 모달(BottomSheet)을 신설하여 쾌적한 UX 완성"),
        ("데이터 객관성 및 신뢰도 논란", "임의로 지정한 '⭐ 인기' 뱃지가 실제 학교 공식 판매량 데이터로 오인될 소지 발견", "임의 수식어 및 인기 태그를 전면 삭제하고, 100% 공식 메뉴명 및 가격표 규격만 순수하게 표시하도록 데이터 무결화 단행"),
        ("실시간 클라우드 서비스 배포 필요", "로컬 환경 실행만으로는 실제 재학생들의 검증 및 시연이 어려움", "Flutter Web Release 파이프라인과 Netlify 호스팅을 연동하여 'https://singular-moonbeam-adafa6.netlify.app'에 24시간 실시간 무중단 배포 완료")
    ]

    for r_idx, r_data in enumerate(t2_rows, start=1):
        row = trouble_table.rows[r_idx]
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.text = val
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            if r_idx % 2 == 1:
                set_cell_background(cell, "F9FAFC")
            cell.paragraphs[0].runs[0].font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 5. 향후 발전 방향 및 결론
    add_h1("5. 기대효과 및 향후 발전 방향 (Conclusion & Future Work)")
    add_h2("5.1 기대 효과")
    add_bullet(" 남서울대 재학생들이 학교 포털, 도서관, 지도, 학식 등 여러 사이트를 전전할 필요 없이 한 곳에서 모든 생활 정보를 실시간 해결", bold_prefix="캠퍼스 정보 접근성 극대화: ")
    add_bullet(" 100% 공식 검증된 가격표와 세트 메뉴를 모바일로 미리 확인하고 키오스크 주문 시간을 단축하여 학생회관 점심 혼잡 완화", bold_prefix="학식 이용 편의 및 혼잡 분산: ")
    add_bullet(" 신입생 및 재학생의 학사 행정 질문을 AI가 24시간 즉시 대응하여 교내 행정 부서의 단순 반복 문의 업무 경감", bold_prefix="행정 효율 향상: ")

    add_h2("5.2 향후 발전 방향")
    add_bullet(" 학교 학생처 및 총학생회와 협의하여 키오스크 POS 실시간 주문 대기 번호 및 품절 현황 API 직접 연계", bold_prefix="키오스크 실시간 주문/품절 API 연동: ")
    add_bullet(" 개인별 시간표 및 수강 과목과 연동한 공결 신청서 자동 생성 및 강의실 이동 최적 경로 안내", bold_prefix="개인화 스마트 시간표 연동: ")
    add_bullet(" 동아리, 학과 학술제, 교내 비교과 마일리지 활동과 포트폴리오 시스템 자동 동기화 기능 확장", bold_prefix="비교과 마일리지 시스템 연계: ")

    add_h2("5.3 최종 결론")
    add_body("본 캡스톤 디자인 프로젝트는 단순한 정보 조회 수준을 넘어, 남서울대학교 재학생들이 일상에서 겪는 실질적인 정보 분절 문제를 AI와 모바일 크로스플랫폼 기술로 완벽히 해결한 실무 중심 프로젝트입니다. 특히 최신 학교 공식 데이터를 100% 동기화하고 최적화된 모달 UI를 구현함으로써 상용 서비스 수준의 완성도를 확보하였으며, 실제 웹으로 배포되어 누구나 즉시 이용 가능한 우수한 성과를 달성하였습니다.")

    # 저장
    import shutil
    out_path1 = r"C:\src\capstone_app\NSU_Capstone_Final_Report.docx"
    out_path2 = r"C:\Users\ohmor\Desktop\NSU_Capstone_Final_Report.docx"
    doc.save(out_path1)
    try:
        shutil.copy2(out_path1, out_path2)
        print(f"Report saved to Desktop: {out_path2}")
    except Exception as e:
        print(f"Could not copy to Desktop: {e}")
    print(f"Report generated successfully: {out_path1}")

if __name__ == "__main__":
    create_report()
